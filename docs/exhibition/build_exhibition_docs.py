from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    bd = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        bd.append(el)
    tc_pr.append(bd)


def add_text(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def build(data, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(2.15)
    sec.right_margin = Cm(2.15)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.8)
    normal.font.color.rgb = RGBColor(35, 40, 45)
    normal.paragraph_format.space_after = Pt(3.5)
    normal.paragraph_format.line_spacing = 1.06
    for sty, size, before, after in (("Title", 19, 0, 9), ("Heading 1", 12.5, 10, 5), ("Heading 2", 10.8, 7, 3)):
        s = styles[sty]
        s.font.name = "Arial"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for sty in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet", "List Number"):
        s = styles[sty]
        s.font.name = "Arial"
        r_pr = s._element.get_or_add_rPr()
        fonts = r_pr.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, fonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), "Arial")
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attr}"), None)
    for el in list(styles["Title"]._element.get_or_add_pPr()):
        if el.tag == qn("w:pBdr"):
            styles["Title"]._element.get_or_add_pPr().remove(el)
    styles["Title"].paragraph_format.line_spacing = 1.05

    p = doc.add_paragraph(style="Title")
    p.add_run(data["title"])
    p = doc.add_paragraph()
    p.add_run(data["subtitle"]).bold = True
    p.paragraph_format.space_after = Pt(12)
    add_text(doc, data["intro"])

    doc.add_heading(data["section_problem"], 1)
    add_text(doc, data["problem_intro"])
    for h, t in data["problems"]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(h + " ").bold = True
        p.add_run(t)
    add_text(doc, data["problem_conclusion"])

    doc.add_heading(data["section_solution"], 1)
    add_text(doc, data["solution_intro"])
    for i, (h, t) in enumerate(data["steps"], 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(f"{h}. ").bold = True
        p.add_run(t)

    doc.add_heading(data["section_scenario"], 1)
    for ptext in data["scenario"]:
        add_text(doc, ptext)

    doc.add_heading(data["section_value"], 1)
    add_text(doc, data["value_intro"])
    for label, value in data["value_rows"]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(label + ". ").bold = True
        p.add_run(value)
    add_text(doc, data["value_note"])

    doc.add_heading(data["section_pitch"], 1)
    add_text(doc, data["pitch"])

    doc.add_heading(data["section_questions"], 1)
    for q, a in data["qa"]:
        doc.add_heading(q, 2)
        add_text(doc, a)

    doc.add_heading(data["section_pilot"], 1)
    add_text(doc, data["pilot_intro"])
    for line in data["pilot_bullets"]:
        add_text(doc, line, "List Bullet")
    add_text(doc, data["pilot_close"])

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(data["footer"]).font.size = Pt(8)
    doc.core_properties.title = data["title"]
    doc.core_properties.subject = data["subject"]
    doc.core_properties.author = "Qural Hub"
    doc.save(path)


from manufacturer_content import RU, KK


if __name__ == "__main__":
    build(RU, OUT / "Qural_Hub_Outage_Impact_RU.docx")
    build(KK, OUT / "Qural_Hub_Outage_Impact_KK.docx")
