from pathlib import Path
import re, json
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(r'C:\Users\BAGDAULET-KOPTILEU\Desktop\AlageumServiceHub\docs\lab7-edit')
SOURCE = Path(r'C:\Users\BAGDAULET-KOPTILEU\Downloads\Kamera-bakylau-tiktok-instagram-target\5655\Lab7_Prolog_Koptileu.docx')
TRANSCRIPT = Path(r'C:\Users\BAGDAULET-KOPTILEU\.codex\attachments\70659655-bbf7-488e-bf96-b97839637971\Вставленный текст.txt')
OUT = ROOT / 'output' / 'Lab7_Prolog_Koptileu_Results.docx'
text = TRANSCRIPT.read_text(encoding='utf-8-sig')
matches = re.findall(r'(?m)^--- TASK (\d+) ---\s*\n(.*?)(?=\n={10,}|\Z)', text, re.S)
results = {}
for number, body in matches:
    number = int(number)
    results.setdefault(number, []).append(f'Select task: {number}\n\n--- TASK {number} ---\n' + body.strip())
assert set(results) == set(range(1, 16)), results.keys()
assert len(results[14]) == 2
doc = Document(SOURCE)
original_code = [p.text for p in doc.paragraphs if p.style.name == 'Program Code']
style = doc.styles.add_style('Program Output', WD_STYLE_TYPE.PARAGRAPH)
style.base_style = doc.styles['Normal']
style.font.name = 'Courier New'
style.font.size = Pt(10)
style.font.color.rgb = RGBColor(0, 0, 0)
style.paragraph_format.line_spacing = Pt(12)
style.paragraph_format.space_before = Pt(2)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.first_line_indent = Pt(0)
style.paragraph_format.keep_together = True
style.paragraph_format.keep_with_next = True
style.paragraph_format.widow_control = True
number = 0
mapping = []
for p in doc.paragraphs:
    if p.text.startswith('Задание '):
        number = int(re.match(r'Задание (\d+)', p.text).group(1))
    if p.text == '[Вставьте скриншот или текст результата выполнения]':
        output = '\n\n'.join(results[number])
        p.clear()
        p.style = style
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(6)
        pf.first_line_indent = Pt(0)
        pf.line_spacing = Pt(12)
        pf.keep_together = True
        pf.keep_with_next = True
        p.add_run(output)
        mapping.append({'task': number, 'runs': len(results[number]), 'text': output})
    elif re.match(r'Рисунок \d+ — Результат выполнения задания \d+$', p.text):
        caption = f'Вывод программы — задание {number}'
        p.clear()
        run = p.add_run(caption)
        run.font.size = Pt(10)
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = False
    elif p.text.startswith('Среда программирования — Visual Prolog 5.0.'):
        p.text = 'Среда программирования — Visual Prolog 5.0. Все 15 заданий объединены в одну программу с текстовым меню. Программный код и результаты выполнения приведены в редактируемом виде.'
    elif p.text.startswith('В разделах «Результат выполнения» следует вставить'):
        p.text = 'В разделах «Результат выполнения» приведены фрагменты предоставленного протокола запуска программы. Повторяющееся главное меню опущено. В задании 14 сохранены обе попытки ввода и полученные ответы программы.'
    elif p.text.startswith('Рекурсивные предикаты выполняют последовательный обход данных'):
        p.text = 'Рекурсивные предикаты выполняют последовательный обход данных, а findall собирает результаты поиска в список. В соответствующих разделах отчёта приведены результаты запусков из предоставленного протокола.'

assert len(mapping) == 15
assert [p.text for p in doc.paragraphs if p.style.name == 'Program Code'] == original_code
assert not any('[Вставьте' in p.text for p in doc.paragraphs)
doc.save(OUT)
check = Document(OUT)
assert [p.text for p in check.paragraphs if p.style.name == 'Program Output'] == [m['text'] for m in mapping]
(ROOT / 'work' / 'mapping.json').write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'replaced': len(mapping), 'runs': sum(m['runs'] for m in mapping), 'code_unchanged': True, 'output': str(OUT)}, ensure_ascii=True))
