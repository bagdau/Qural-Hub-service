from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "docs/presentation/deck-output/Qural_Hub_3_Minutes.pptx"
ASSETS = Path(__file__).resolve().parent / "assets"
LOGO = ROOT / "docs/presentation/deck-work/university-logo.png"

NAVY = RGBColor(24, 35, 65)
PURPLE = RGBColor(88, 65, 222)
LIGHT = RGBColor(245, 246, 251)
TEXT = RGBColor(25, 36, 62)
MUTED = RGBColor(90, 100, 119)
WHITE = RGBColor(255, 255, 255)


def box(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def text(slide, value, x, y, w, h, size, color=TEXT, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.text = value
    if align is not None:
        p.alignment = align
    p.font.name = "Arial"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def footer(slide, num, dark=False):
    c = RGBColor(194, 201, 222) if dark else MUTED
    text(slide, "Qural Hub · демонстрационный сценарий", .6, 7.12, 7.2, .18, 9, c)
    text(slide, f"0{num}", 12.35, 7.1, .38, .2, 9, c, align=PP_ALIGN.RIGHT)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# 1. Cover: producer's equipment and its service record.
s = prs.slides.add_slide(blank)
box(s, 0, 0, 13.333, 7.5, NAVY)
if LOGO.exists():
    s.shapes.add_picture(str(LOGO), Inches(.65), Inches(.43), width=Inches(1.35))
text(s, "Qural Hub", .65, 1.65, 5.1, .73, 42, WHITE, True)
text(s, "Сервис производителя\nтрансформаторов и КТП", .68, 2.52, 5.0, 1.4, 27, WHITE, True)
text(s, "От ввода в эксплуатацию и гарантии\nдо заявки, ремонта и замены", .68, 4.24, 4.8, 1.1, 18, RGBColor(208, 210, 242))
s.shapes.add_picture(str(ASSETS / "equipment-2026-10-27.png"), Inches(6.0), Inches(1.5), width=Inches(6.75), height=Inches(3.8))
text(s, "Реальный экран Qural Hub · демонстрационные данные", 6.02, 5.55, 6.6, .33, 11, RGBColor(208, 210, 242))
footer(s, 1, True)

# 2. Impact: one supplying node may affect several consumers.
s = prs.slides.add_slide(blank)
box(s, 0, 0, 13.333, 7.5, WHITE)
text(s, "Один сбой питания — много последствий", .6, .42, 12.1, .55, 27, TEXT, True)
text(s, "Если потребители питаются от одной КТП, её отказ затрагивает их одновременно.", .62, 1.08, 12, .4, 16, MUTED)
box(s, .62, 1.69, 12.05, .03, PURPLE)
for y, title, body in [
    (1.94, "Продукты", "Холодильники и морозильники останавливаются. При долгом отключении мясо, молочные продукты и мороженое могут испортиться."),
    (3.19, "Жара и медицина", "Летом перестают работать кондиционеры. Охлаждение и вентиляция медучреждений зависят от возможностей резервного питания."),
    (4.44, "Производство и офис", "Станки, насосы и подъёмные механизмы простаивают; без достаточного резерва останавливаются серверы, компьютеры и кассы."),
]:
    text(s, title, .65, y, 3.2, .38, 22, PURPLE, True)
    text(s, body, 3.85, y-.02, 8.72, .9, 18, TEXT)
    if y < 4:
        box(s, .65, y+1.07, 11.9, .012, RGBColor(220, 223, 233))
text(s, "Qural Hub помогает быстрее найти объект, проверить обязательства и организовать выезд.", .65, 6.36, 12, .38, 17, TEXT, True)
footer(s, 2)

# 3. Response: real newly captured board screenshot with the proposed workflow.
s = prs.slides.add_slide(blank)
box(s, 0, 0, 13.333, 7.5, WHITE)
text(s, "От обращения до выезда", .6, .42, 12.1, .55, 27, TEXT, True)
text(s, "Qural Hub связывает отключение с установленным оборудованием и историей его обслуживания.", .62, 1.08, 12, .45, 16, MUTED)
items = [
    ("01", "Сообщение", "Житель называет адрес и наблюдаемый факт. При угрозе людям — 112."),
    ("02", "Проверка", "Диспетчер связывает обращения, находит паспорт КТП и условия гарантии."),
    ("03", "Выезд", "Квалифицированные специалисты устанавливают причину и выполняют работы."),
    ("04", "Результат", "Владелец видит статус и акт; после срока готовится обоснованная замена."),
]
for i, (n, title, body) in enumerate(items):
    y = 1.79 + i * 1.05
    text(s, n, .65, y, .52, .38, 18, PURPLE, True)
    text(s, title, 1.27, y-.02, 4.7, .3, 17, TEXT, True)
    text(s, body, 1.27, y+.30, 4.75, .65, 13.5, MUTED)
s.shapes.add_picture(str(ASSETS / "board-2026-10-27.png"), Inches(6.25), Inches(1.74), width=Inches(6.48), height=Inches(3.65))
text(s, "Новый снимок экрана · демонстрационные заявки", 6.26, 5.55, 6.4, .25, 10, MUTED)
box(s, .64, 6.30, 12.0, .02, RGBColor(220, 223, 233))
text(s, "Сегодня: паспорта и заявки. Связь с 109, датчики и назначение бригад — следующий этап пилота.", .65, 6.48, 12.0, .42, 14, TEXT, True)
footer(s, 3)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
