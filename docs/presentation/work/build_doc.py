from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from content import PAGES

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "Alageum_Service_Hub_Overview_Technical.docx"
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin,sec.bottom_margin = Cm(1.7),Cm(1.6)
sec.left_margin,sec.right_margin = Cm(1.85),Cm(1.85)
sec.header_distance, sec.footer_distance = Cm(.65),Cm(.7)
sec.different_first_page_header_footer = True
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    styles[name].font.name='Arial'
    styles[name].font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(12)
styles['Normal'].paragraph_format.line_spacing=1.15
styles['Normal'].paragraph_format.space_after=Pt(8)
styles['Title'].font.size=Pt(32)
styles['Title'].font.bold=True
styles['Subtitle'].font.size=Pt(19)
styles['Subtitle'].font.italic=False
styles['Heading 1'].font.size=Pt(21)
styles['Heading 1'].font.bold=True
styles['Heading 1'].paragraph_format.space_after=Pt(12)
styles['Heading 1'].paragraph_format.space_before=Pt(0)
styles['Heading 2'].font.size=Pt(12)
styles['Heading 2'].font.bold=True
styles['Caption'].font.size=Pt(9)
styles['Caption'].font.italic=False
styles['Caption'].paragraph_format.space_after=Pt(10)
styles['Caption'].paragraph_format.line_spacing=1.05
for st in ['Title','Subtitle','Heading 1','Heading 2']:
    styles[st].paragraph_format.keep_with_next=True
for st in styles:
    for borders in st.element.xpath('.//w:pBdr'):
        borders.getparent().remove(borders)
    for color in st.element.xpath('.//w:color'):
        for a in ['themeColor','themeTint','themeShade']:
            color.attrib.pop(qn('w:'+a),None)

header=sec.header.paragraphs[0]
header.text="Qural Hub    /    ОБЗОР И ТЕХНИЧЕСКОЕ ОПИСАНИЕ"
header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Qural Hub   •   ").font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
footer._p.append(field)
doc.core_properties.title="Qural Hub"
doc.core_properties.subject="Презентационное и техническое описание программы"
doc.core_properties.author="Команда проекта Qural Hub"
doc.core_properties.keywords="сервис, оборудование, заявки, SLA, руководство"

def para(text,style=None):
    p=doc.add_paragraph(text,style)
    p.paragraph_format.widow_control=True
    return p

def bookmark(p,name,n):
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(n));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(n))
    p._p.insert(0,a);p._p.append(b)

def picture(name,width=17.3,height=None):
    path=ASSETS/(name+'.png')
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_with_next=True
    if height:shape=p.add_run().add_picture(str(path),height=Cm(height))
    else:shape=p.add_run().add_picture(str(path),width=Cm(width))
    shape._inline.docPr.set('descr',name.replace('-',' ')+" — экран Qural Hub")
    return p

def table(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    ratios=[.25,.75] if len(headers)==2 else [.25,.26,.49]
    if headers[0]=="Роль":ratios=[.18,.45,.37]
    if headers[0]=="Этап":ratios=[.29,.24,.47]
    if headers[0]=="Параметр":ratios=[.43,.57]
    if headers[0].startswith("Маршрут"):ratios=[.50,.12,.38]
    if headers==["Приоритет","Реакция","Выполнение"]:ratios=[.4,.3,.3]
    for c,r in zip(t.columns,ratios):c.width=Cm(17.3*r)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=text
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,c in enumerate(row.cells):
            c.width=Cm(17.3*ratios[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'263E60' if ri==0 else ('F1F4F8' if ri%2==0 else 'FFFFFF'));tcpr.append(shd)
            margins=OxmlElement('w:tcMar')
            for name in ['top','left','bottom','right']:
                x=OxmlElement('w:'+name);x.set(qn('w:w'),'85');x.set(qn('w:type'),'dxa');margins.append(x)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for name in ['top','left','bottom','right']:
                x=OxmlElement('w:'+name);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
            tcpr.append(borders)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1)
                p.paragraph_format.space_before=Pt(1)
                p.paragraph_format.line_spacing=1.04
                for run in p.runs:
                    run.font.size=Pt(10);run.bold=ri==0
                    run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
    para('').paragraph_format.space_after=Pt(0)

# Cover
para("Qural Hub","Title")
para("Презентационное и техническое\nописание программы","Subtitle")
para("Учёт электротехнического оборудования\nи управление сервисным обслуживанием")
picture('dashboard')
para("Функции и интерфейс • Процесс обслуживания • Архитектура и эксплуатация","Caption")
para("Демонстрационная версия MVP")
para("20 октября 2026 года")

# Navigation
doc.add_page_break()
para("Содержание","Heading 1")
para("Разделы 1–25 знакомят с продуктом и рабочими сценариями. Разделы 26–34 описывают устройство, сопровождение и приёмку системы.")
for i,item in enumerate(PAGES,1):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.line_spacing=1
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.9),WD_ALIGN_PARAGRAPH.RIGHT)
    hl=OxmlElement('w:hyperlink');hl.set(qn('w:anchor'),f'chapter{i}')
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz);r.append(pr)
    tx=OxmlElement('w:t');tx.text=f'{i:02d}  {item["title"]}';r.append(tx);hl.append(r);p._p.append(hl)
    p.add_run(f'\t{i+2}').font.size=Pt(10)

fig=0
for idx,item in enumerate(PAGES,1):
    doc.add_page_break()
    h=para(f'{idx:02d}  {item["title"]}',"Heading 1")
    bookmark(h,f'chapter{idx}',idx)
    if item.get('image'):
        fig+=1
        # a compact screen leaves room for description and the policy table
        picture(item['image'],width=15.6 if item['image']=='templates' else 17.3)
        para(f'Рисунок {fig}. {item["caption"]}',"Caption")
    if item.get('images'):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        for j,name in enumerate(item['images']):
            if j:p.add_run('     ')
            sh=p.add_run().add_picture(str(ASSETS/(name+'.png')),height=Cm(12.2))
            sh._inline.docPr.set('descr',name+" мобильный интерфейс")
        fig+=1;para(f'Рисунок {fig}. {item["caption"]}',"Caption")
    for text in item.get('paras',[]):para(text)
    if item.get('table'):table(*item['table'])
    for text in item.get('after',[]):para(text)
    p=para('Основа описания: '+item['source'])
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(0)
    for r in p.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('555555')

doc.save(OUTPUT)
print(f'Created {OUTPUT}; intended pages={len(PAGES)+2}; figures={fig}')
